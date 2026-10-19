#!/usr/bin/env python3
"""Skrip untuk mengisi data Jabatan, Kelas Jabatan, dan Masa Kerja dari Excel nominatif ke database SICUTI."""
import os
import sys
import re
import shutil
import sqlite3
import argparse
import openpyxl
from werkzeug.security import generate_password_hash

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "database.db")
BACKUP_PATH = os.path.join(BASE_DIR, "database_sebelum_nominatif.db")

def hitung_tmt(acuan_str, masa_kerja_str):
    if not masa_kerja_str or not str(masa_kerja_str).strip():
        return ""
    m = re.search(r'(\d+)\s*Tahun\s*(\d+)\s*Bulan', str(masa_kerja_str), re.IGNORECASE)
    if not m:
        return ""
    th = int(m.group(1))
    bl = int(m.group(2))
    selisih_bulan = th * 12 + bl
    try:
        acuan_th, acuan_bl = map(int, acuan_str.strip()[:7].split("-"))
    except Exception:
        acuan_th, acuan_bl = 2026, 10
    total_acuan_m = acuan_th * 12 + (acuan_bl - 1)
    target_m = total_acuan_m - selisih_bulan
    target_th = target_m // 12
    target_bl = (target_m % 12) + 1
    return f"{target_th:04d}-{target_bl:02d}-01"

def main():
    parser = argparse.ArgumentParser(description="Isi data nominatif ke database SICUTI")
    parser.add_argument("xlsx", nargs="?", default="nominatif.xlsx", help="Path file nominatif.xlsx")
    parser.add_argument("--acuan", default="2026-10", help="Bulan acuan YYYY-MM (default: 2026-10)")
    parser.add_argument("--tambah-baru", action="store_true", help="Tambahkan pegawai yang belum ada di database")
    args = parser.parse_args()

    xlsx_path = args.xlsx
    if not os.path.isabs(xlsx_path):
        xlsx_path = os.path.join(BASE_DIR, xlsx_path)

    if not os.path.exists(xlsx_path):
        print(f"Error: File Excel {xlsx_path} tidak ditemukan.")
        sys.exit(1)

    sys.path.insert(0, BASE_DIR)
    from db_config import get_db, DB_TYPE
    from app import init_db

    # 1. Cadangkan database jika SQLite
    if DB_TYPE == "sqlite" and os.path.exists(DB_PATH):
        shutil.copy2(DB_PATH, BACKUP_PATH)
        print(f"Cadangan database berhasil dibuat: {os.path.basename(BACKUP_PATH)}")

    # 2. Inisialisasi / migrasi kolom database lewat init_db app.py
    init_db()

    # Buka database
    con = get_db()
    cur = con.cursor()

    r_th = cur.execute("SELECT nilai FROM pengaturan WHERE kunci='tahun_aktif'").fetchone()
    th_aktif = int(r_th["nilai"]) if r_th and r_th["nilai"] else 2026

    # Ambil data pegawai role 'pegawai' yang ada di DB
    db_pegawai = {}
    for p in cur.execute("SELECT * FROM pegawai WHERE role='pegawai'").fetchall():
        if p["nip"]:
            digits = re.sub(r"\D", "", str(p["nip"]))
            if digits:
                db_pegawai[digits] = p

    # 3. Baca Excel nominatif
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet = wb.active

    # Deteksi baris judul (cari sel "NIP")
    header_row = None
    col_idx = {}
    for r in range(1, 15):
        for c in range(1, sheet.max_column + 1):
            val = sheet.cell(r, c).value
            if val and "NIP" in str(val).strip().upper():
                header_row = r
                break
        if header_row:
            for c in range(1, sheet.max_column + 1):
                hval = sheet.cell(header_row, c).value
                if hval:
                    col_idx[str(hval).strip().upper()] = c
            break

    if not header_row or "NIP" not in col_idx:
        print("Error: Tidak dapat menemukan baris judul dengan kolom NIP di file Excel.")
        con.close()
        sys.exit(1)

    c_unor = col_idx.get("UNOR")
    c_nip = col_idx.get("NIP")
    c_nama = col_idx.get("NAMA")
    c_gol = col_idx.get("GOLONGAN")
    c_jabatan = col_idx.get("JABATAN")
    c_mk = col_idx.get("MASA KERJA")
    c_kj = col_idx.get("KELAS JABATAN")


    diperbarui_count = 0
    ditambah_count = 0
    dilewati = []
    tanpa_masa_kerja = []
    kadin_info = None

    for r in range(header_row + 1, sheet.max_row + 1):
        nip_val = sheet.cell(r, c_nip).value if c_nip else None
        if not nip_val:
            continue
        nip_digits = re.sub(r"\D", "", str(nip_val))
        if not nip_digits:
            continue

        nama_val = str(sheet.cell(r, c_nama).value or "").strip() if c_nama else ""
        unor_val = str(sheet.cell(r, c_unor).value or "").strip() if c_unor else ""
        gol_val = str(sheet.cell(r, c_gol).value or "").strip() if c_gol else ""
        jabatan_val = str(sheet.cell(r, c_jabatan).value or "").strip() if c_jabatan else ""
        mk_raw = sheet.cell(r, c_mk).value if c_mk else ""
        mk_val = str(mk_raw).strip() if mk_raw is not None else ""
        kj_raw = sheet.cell(r, c_kj).value if c_kj else ""
        kj_val = str(kj_raw).strip() if kj_raw is not None else ""

        # Deteksi Kepala Dinas Perhubungan
        if jabatan_val.lower() == "kepala dinas perhubungan":
            cur.execute("INSERT OR REPLACE INTO pengaturan VALUES('kepala_nama', 'CHRESTINA ERNI WIDYASTUTI, S.E., M.Si.')")
            cur.execute("INSERT OR REPLACE INTO pengaturan VALUES('kepala_nip', '196905021991022002')")
            kadin_info = ("CHRESTINA ERNI WIDYASTUTI, S.E., M.Si.", "196905021991022002")

        # Hitung TMT
        tmt_val = hitung_tmt(args.acuan, mk_val)

        # Cek apakah pegawai ada di DB
        if nip_digits in db_pegawai:
            peg_id = db_pegawai[nip_digits]["id"]
            if not mk_val or not tmt_val:
                tanpa_masa_kerja.append((nip_digits, nama_val or db_pegawai[nip_digits]["nama"]))
                tmt_val = ""

            cur.execute("""UPDATE pegawai SET jabatan=?, kelas_jabatan=?, masa_kerja=?, tmt_kerja=?
                WHERE id=?""", (jabatan_val, kj_val, mk_val, tmt_val, peg_id))
            diperbarui_count += 1
        else:
            if args.tambah_baru:
                if not mk_val or not tmt_val:
                    tanpa_masa_kerja.append((nip_digits, nama_val))
                    tmt_val = ""
                cur.execute("""INSERT INTO pegawai(nip, nama, golongan, unit_kerja, username,
                    password_hash, role, aktif, jabatan, kelas_jabatan, masa_kerja, tmt_kerja)
                    VALUES(?, ?, ?, ?, ?, ?, 'pegawai', 1, ?, ?, ?, ?)""",
                    (nip_digits, nama_val, gol_val, unor_val, nip_digits,
                     generate_password_hash(nip_digits), jabatan_val, kj_val, mk_val, tmt_val))
                new_pid = cur.lastrowid
                cur.execute("""INSERT OR IGNORE INTO jatah_cuti(pegawai_id, tahun, jatah_dasar, total_jatah, sisa_cuti)
                    VALUES(?, ?, 12, 12, 12)""", (new_pid, th_aktif))
                ditambah_count += 1
            else:
                dilewati.append((nip_digits, nama_val, jabatan_val))

    con.commit()
    con.close()

    # Cetak laporan akhir
    print("=" * 60)
    print("LAPORAN PENGISIAN DATA NOMINATIF")
    print("=" * 60)
    print(f"Hasil: {diperbarui_count} pegawai diperbarui")
    if args.tambah_baru:
        print(f"Hasil: {ditambah_count} pegawai baru ditambahkan")

    if kadin_info:
        print(f"Kepala Dinas diperbarui: {kadin_info[0]} ({kadin_info[1]})")

    print(f"\nNIP di Excel yang belum ada di database (dilewati): {len(dilewati)}")
    for nip, nama, jab in dilewati:
        print(f"  - NIP: {nip} | Nama: {nama} | Jabatan: {jab}")

    print(f"\nPegawai tanpa masa kerja (PPPK): {len(tanpa_masa_kerja)}")
    for nip, nama in tanpa_masa_kerja:
        print(f"  - NIP: {nip} | Nama: {nama}")
    print("=" * 60)

if __name__ == "__main__":
    main()
